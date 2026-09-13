import io
import logging
from typing import Optional
from fastapi import HTTPException, status
import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

logger = logging.getLogger("careerx.text_extraction")


def extract_text_from_pdf(content: bytes) -> str:
    """Extract parseable text from raw PDF bytes using pypdf.
    
    Raises:
        HTTPException 400 on encryption, corruption, empty file, or unparseable scanned PDF.
    """
    if not content or len(content) < 4:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="PDF file is empty or corrupted.",
        )

    try:
        stream = io.BytesIO(content)
        reader = PdfReader(stream)
    except Exception as e:
        logger.warning("Failed to parse PDF stream: %s", e)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded PDF file is corrupted or could not be read.",
        )

    if reader.is_encrypted:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded PDF is password-protected or encrypted. Please upload an unprotected PDF.",
        )

    if len(reader.pages) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded PDF contains zero pages.",
        )

    extracted_pages = []
    for page_idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                extracted_pages.append(page_text.strip())
        except Exception as e:
            logger.warning("Error extracting text from PDF page %d: %s", page_idx, e)

    combined_text = "\n\n".join(extracted_pages).strip()

    # Rejection if no extractable text found (e.g. pure raster image scan)
    if not combined_text or len(combined_text) < 20:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "No extractable text found in PDF. Scanned image PDFs without OCR "
                "are not supported. Please upload a PDF containing selectable text or a DOCX document."
            ),
        )

    return combined_text


def extract_text_from_docx(content: bytes) -> str:
    """Extract parseable text from raw DOCX bytes using python-docx.
    
    Raises:
        HTTPException 400 on corruption, empty file, or unparseable document.
    """
    if not content or len(content) < 4:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="DOCX file is empty or corrupted.",
        )

    try:
        stream = io.BytesIO(content)
        doc = docx.Document(stream)
    except Exception as e:
        logger.warning("Failed to parse DOCX stream: %s", e)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded DOCX file is corrupted or not a valid Word document.",
        )

    parts = []
    # 1. Paragraphs
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if p_text:
            parts.append(p_text)

    # 2. Tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))

    combined_text = "\n\n".join(parts).strip()

    if not combined_text or len(combined_text) < 20:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No extractable text found in DOCX file. Please upload a document containing resume text.",
        )

    return combined_text


def extract_resume_text(content: bytes, file_format: str) -> str:
    """Dispatches text extraction based on file format ('PDF' or 'DOCX')."""
    fmt = file_format.upper().strip()
    if fmt == "PDF":
        return extract_text_from_pdf(content)
    elif fmt == "DOCX":
        return extract_text_from_docx(content)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported resume format '{file_format}'. Allowed formats: PDF, DOCX.",
        )
